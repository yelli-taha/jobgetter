"""Build and check application documents for one job.

Usage (from the project folder; PY = .venv\\Scripts\\python.exe on Windows, .venv/bin/python on Mac):
  PY tools/build.py resume <job_id>
  PY tools/build.py letter <job_id>
  PY tools/build.py all <job_id>
  PY tools/build.py check <file.tex>      check any resume .tex without a job folder
  PY tools/build.py master                compile me/samples/master-resume.tex -> master-resume.pdf

Inputs in jobs/<job_id>/:
  resume.tex         tailored copy of me/samples/master-resume.tex (written by Claude with tools/resume_kit.py)
  cover-letter.txt   letter text: salutation line ... sign-off line, then the user's name (see me/profile.json)

Outputs: named PDFs (+ Word letter), and mechanical rule checks saved into job.json.
The rules come from me/profile.json and me/checks.json. Claude still does the judgement checks
(voice, relevance, facts) before showing a draft.
"""
import json
import re
import shutil
import subprocess
import sys
import tempfile
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
from pypdf import PdfReader  # noqa: E402

from tools.common import JOBS, ME, checks as load_checks, profile as load_profile  # noqa: E402

MASTER = ME / "samples" / "master-resume.tex"
SKILLS_ROWS = r"\\\\\[1pt\]\n(.*?)\n&\n(.*?)\\\\\n\\end\{tabularx\}"


# ---------- helpers ----------

def pdflatex():
    found = shutil.which("pdflatex")
    if found:
        return found
    home = Path.home()
    candidates = [home / "AppData/Local/Programs/MiKTeX/miktex/bin/x64/pdflatex.exe",
                  Path("C:/Program Files/MiKTeX/miktex/bin/x64/pdflatex.exe"),
                  Path("/Library/TeX/texbin/pdflatex"), Path("/usr/local/bin/pdflatex"), Path("/opt/homebrew/bin/pdflatex")]
    candidates += sorted(home.glob("Library/TinyTeX/bin/*/pdflatex")) + sorted(home.glob(".TinyTeX/bin/*/pdflatex"))
    candidates += sorted(home.glob("AppData/Roaming/TinyTeX/bin/*/pdflatex.exe"))
    for p in candidates:
        if p.exists():
            return str(p)
    sys.exit("pdflatex not found. Is MiKTeX (Windows) or TinyTeX/MacTeX (Mac) installed?")


def compile_tex(tex_path: Path, out_pdf: Path):
    """Compile in a temp folder so log/aux files don't clutter the job folder."""
    with tempfile.TemporaryDirectory() as tmp:
        shutil.copy(tex_path, Path(tmp) / "doc.tex")
        for _ in range(2):  # twice, so references settle
            r = subprocess.run([pdflatex(), "-interaction=nonstopmode", "-halt-on-error", "doc.tex"],
                               cwd=tmp, capture_output=True, text=True, errors="replace")
        pdf = Path(tmp) / "doc.pdf"
        if r.returncode != 0 or not pdf.exists():
            errors = [l for l in r.stdout.splitlines() if l.startswith("!")]
            raise RuntimeError("LaTeX failed: " + ("; ".join(errors[:3]) or r.stdout[-500:]))
        shutil.copy(pdf, out_pdf)
    return len(PdfReader(str(out_pdf)).pages)


def slug(text):
    return re.sub(r"[^A-Za-z0-9]+", "_", text).strip("_")


def load_job(job_id):
    path = JOBS / job_id / "job.json"
    return path, json.loads(path.read_text(encoding="utf-8"))


def latex_escape(text):
    repl = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
            "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}
    return "".join(repl.get(c, c) for c in text)


def issue(level, text):
    return {"level": level, "text": text}


def block(text, start, end):
    m = re.search(re.escape(start) + r"(.*?)" + re.escape(end), text, flags=re.S)
    return m.group(1) if m else None


def sections(body):
    return re.findall(r"\\section\{(.*?)\}", body)


# ---------- resume ----------

def resume_items(tex):
    return re.findall(r"\\resumeItem\{(.*)\}\s*$", tex, flags=re.M)


def check_skills(body, master_body, ck):
    """Skills table (two cells): layout must match the master; words checked against me/checks.json lists."""
    out = []

    def table(text):
        r = re.search(SKILLS_ROWS, text, flags=re.S)
        if not r:
            return None, None
        start = text.rfind(r"\begin{tabularx}", 0, r.start(1))
        t = text[start:text.find(r"\end{tabularx}", r.end(2)) + len(r"\end{tabularx}")]
        return t, re.search(SKILLS_ROWS, t, flags=re.S)

    t_master, m_rows = table(master_body)
    if not m_rows:
        return out  # the master has no two-cell skills table: nothing to check
    t_now, rows = table(body)
    if not rows:
        return [issue("error", "Skills table layout changed (couldn't find its two cells)")]
    frame = lambda s, r: s[:r.start(1)] + s[r.end(1):r.start(2)] + s[r.end(2):]
    if frame(t_now, rows) != frame(t_master, m_rows):
        out.append(issue("error", "Skills table layout differs from the master (only the words inside may change)"))
    unesc = lambda s: s.replace(r"\&", "&").strip()
    left = [unesc(t) for t in rows.group(1).split(r"$\cdot$") if t.strip()]
    right = [unesc(c) for c in rows.group(2).split(r"\newline") if c.strip()]
    if ck["skills_left"]:
        for t in left:
            if t not in ck["skills_left"]:
                out.append(issue("error", f"Not in your skills list (or renamed): \"{t}\""))
    allowed = ck["skills_right_default"] + ck["skills_right_stockpile"]
    for c in right:
        if any(c.lower() == r.lower() for r in ck["skills_right_rejected"]):
            out.append(issue("error", f"Rejected wording: \"{c}\""))
        elif allowed and c not in allowed:
            out.append(issue("warn", f"\"{c}\" isn't in your skills list. Make sure it's true and explain why under Skills chosen."))
        if ck["max_skill_chars"] and len(c) > ck["max_skill_chars"]:
            out.append(issue("error", f"Too long for one line: \"{c}\""))
    missing = [d for d in ck["skills_right_default"] if d not in right]
    if missing:
        out.append(issue("error", "Always-included skills missing: " + ", ".join(missing)))
    if len(ck["skills_right_count"]) == 2 and not ck["skills_right_count"][0] <= len(right) <= ck["skills_right_count"][1]:
        out.append(issue("error", f"{len(right)} items in the right-hand skills column "
                                  f"(must be {ck['skills_right_count'][0]} to {ck['skills_right_count'][1]})"))
    if not out:
        out.append(issue("ok", f"Skills table intact: {len(left)} + {len(right)} items"))
    return out


def check_resume(tex, pages):
    pr, ck = load_profile(), load_checks()
    master = MASTER.read_text(encoding="utf-8")
    m_body = master.split(r"\begin{document}", 1)[-1]
    body = tex.split(r"\begin{document}", 1)[-1]
    max_pages = pr["resume"].get("max_pages", 1)
    out = [issue("ok" if pages <= max_pages else "error", f"Length: {pages} page{'s' if pages != 1 else ''} (max {max_pages})")]
    # Section order: same order as the master for the sections both have
    m_secs, secs = sections(m_body), sections(body)
    shared_master = [s for s in m_secs if s in secs]
    shared_now = [s for s in secs if s in m_secs]
    if shared_master != shared_now:
        out.append(issue("error", "Sections are in a different order from the master: " + ", ".join(secs)))
    out += check_skills(body, m_body, ck)
    for pat, why in ck["resume_banned_patterns"].items():
        if re.search(pat, body, flags=re.I):
            out.append(issue("error", why))
    for rule in ck["never_together"]:
        if rule.get("where", "both") in ("resume", "both"):
            m = re.search(r"% JOB: " + re.escape(rule.get("job_key", "")) + r"\b.*?\\resumeItemListEnd", body, flags=re.S)
            if m and rule["b"].lower() in m.group(0).lower():
                out.append(issue("error", rule.get("why") or f"{rule['b']} appears under {rule['a']}"))
    if tex.split(r"\begin{document}", 1)[0] != master.split(r"\begin{document}", 1)[0]:
        out.append(issue("error", "Layout/preamble differs from the master (only content may change)"))
    date_pat = r"\b(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Sept|Oct|Nov|Dec)[a-z]*\.? \d{4}"
    new_dates = set(re.findall(date_pat, body)) - set(re.findall(date_pat, master))
    for d in sorted(new_dates):
        out.append(issue("error", f"Date not in the master resume: {d}"))
    for item in resume_items(body):
        if re.search(r"(?<!\\)[&%$#_]", item):
            out.append(issue("error", f"Unescaped special character in: {item[:60]}..."))
    facts = master + "".join(p.read_text(encoding="utf-8") for p in (ME / "experience").glob("*.md"))
    known_nums = set(re.findall(r"\d+(?:\.\d+)?", facts))
    for item in resume_items(body):
        for n in re.findall(r"\d+(?:\.\d+)?", item):
            if n not in known_nums:
                out.append(issue("warn", f"Number {n} isn't in your experience files. Check it: \"{item[:70]}...\""))
    if not any(i["level"] in ("error", "warn") for i in out):
        out.append(issue("ok", "No mechanical rule problems found"))
    return out


def build_resume_checks_only(tex_path):
    with tempfile.TemporaryDirectory() as tmp:
        pages = compile_tex(Path(tex_path), Path(tmp) / "x.pdf")
    return check_resume(Path(tex_path).read_text(encoding="utf-8"), pages)


def build_resume(job_id):
    path, job = load_job(job_id)
    folder = JOBS / job_id
    tex_src = folder / "resume.tex"
    if not tex_src.exists():
        sys.exit(f"Missing {tex_src}")
    name = f"{load_profile()['file_prefix']}_Resume_{slug(job['company'])}_{slug(job['title'])}"
    shutil.copy(tex_src, folder / f"{name}.tex")
    pages = compile_tex(tex_src, folder / f"{name}.pdf")
    checks = check_resume(tex_src.read_text(encoding="utf-8"), pages)
    d = job.setdefault("drafts", {})
    d.update({"resume_tex": f"{job_id}/{name}.tex", "resume_pdf": f"{job_id}/{name}.pdf",
              "resume_is_placeholder": False, "resume_checks": checks})
    d.setdefault("resume_version", "main")  # see me/samples/resume-versions/versions.md
    path.write_text(json.dumps(job, indent=2, ensure_ascii=False), encoding="utf-8")
    return checks


# ---------- cover letter ----------

LETTER_TEMPLATE = r"""\documentclass[<<PAPER>>,11pt]{article}
\usepackage[utf8]{inputenc}
\input{glyphtounicode}
\pdfgentounicode=1
\usepackage[margin=1in]{geometry}
\usepackage[hidelinks]{hyperref}
\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{9pt}
\begin{document}
\begin{center}
  {\Huge \scshape <<NAME>>} \\ \vspace{1pt}
  <<CONTACT>>
\end{center}
\vspace{4pt}
<<DATE>>

<<BODY>>
\end{document}
"""


def contact_bits(pr):
    return [x for x in (pr["email"], pr["phone"], re.sub(r"^https?://(www\.)?", "", pr["linkedin"] or "").rstrip("/"),
                        re.sub(r"^https?://(www\.)?", "", pr["website"] or "").rstrip("/")) if x]


def contact_tex(pr):
    parts = []
    for x in contact_bits(pr):
        if "@" in x:
            parts.append(r"\href{mailto:" + x + "}{" + latex_escape(x) + "}")
        elif "." in x and "/" in x:
            parts.append(r"\href{https://" + x + "}{" + latex_escape(x) + "}")
        else:
            parts.append(latex_escape(x))
    return " $|$ ".join(parts)


def letter_paragraphs(text):
    return [p.strip() for p in re.split(r"\n\s*\n", text.strip()) if p.strip()]


def check_letter(text, pages):
    pr, ck = load_profile(), load_checks()
    lt = pr["letter"]
    out = [issue("ok" if pages <= lt.get("max_pages", 1) else "error",
                 f"Length: {pages} page{'s' if pages != 1 else ''} (max {lt.get('max_pages', 1)})")]
    paras = letter_paragraphs(text)
    body = "\n\n".join(paras[1:-1]) if len(paras) > 2 else text
    words = len(re.findall(r"[A-Za-z0-9'’]+", body))
    lo, hi = lt.get("min_words", 250), lt.get("max_words", 500)
    out.append(issue("ok" if lo <= words <= hi else "warn", f"{words} words in the body (aim for {lo} to {hi})"))
    if lt.get("salutation") and (not paras or paras[0] != lt["salutation"]):
        out.append(issue("error", f"Must open with exactly \"{lt['salutation']}\""))
    ending = f"{lt.get('signoff', 'Sincerely,')}\n{pr['name']}"
    if pr["name"] and not text.rstrip().endswith(ending):
        out.append(issue("error", f"Must end with \"{lt.get('signoff', 'Sincerely,')}\" then \"{pr['name']}\""))
    if ck["no_em_dashes"] and ("—" in text or "--" in text or re.search(r"\s–\s", text)):
        out.append(issue("error", "Contains an em dash (or a dash used like one)"))
    low = text.lower().replace("’", "'")
    for p in ck["letter_banned_phrases"]:
        if p.lower() in low:
            out.append(issue("error", f"Banned phrase: \"{p}\""))
    for pat, why in ck["letter_banned_patterns"].items():
        if re.search(pat, text, flags=re.I):
            out.append(issue("error", why))
    for rule in ck["never_together"]:
        if rule.get("where", "both") in ("letter", "both"):
            for sent in re.split(r"(?<=[.!?])\s+", text):
                if rule["a"].lower() in sent.lower() and rule["b"].lower() in sent.lower():
                    out.append(issue("error", rule.get("why") or f"{rule['b']} mentioned with {rule['a']}"))
    if ck["no_paragraph_headers"]:
        for p in paras[1:-1]:
            if len(p.split()) <= 6 and not p.rstrip().endswith((".", "!", "?", ":")):
                out.append(issue("error", f"Looks like a paragraph header: \"{p}\""))
    if not any(i["level"] in ("error", "warn") for i in out):
        out.append(issue("ok", "No mechanical rule problems found"))
    return out


def build_letter(job_id):
    pr = load_profile()
    path, job = load_job(job_id)
    folder = JOBS / job_id
    text = (folder / "cover-letter.txt").read_text(encoding="utf-8").replace("\r\n", "\n").strip()
    paras = letter_paragraphs(text)
    d0 = date.today()
    today = f"{d0.strftime('%B')} {d0.day}, {d0.year}"

    tex_paras = [latex_escape(p).replace("\n", r"\\" + "\n") for p in paras]
    tex = (LETTER_TEMPLATE.replace("<<PAPER>>", pr.get("paper", "a4paper"))
           .replace("<<NAME>>", latex_escape(pr["name"])).replace("<<CONTACT>>", contact_tex(pr))
           .replace("<<DATE>>", today).replace("<<BODY>>", "\n\n".join(tex_paras)))
    name = f"{pr['file_prefix']}_Cover_Letter_{slug(job['company'])}_{slug(job['title'])}"
    tex_path = folder / f"{name}.tex"
    tex_path.write_text(tex, encoding="utf-8")
    pages = compile_tex(tex_path, folder / f"{name}.pdf")
    tex_path.unlink()  # the editable versions are cover-letter.txt and the Word file

    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt
    doc = Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = s.left_margin = s.right_margin = Inches(1)
    style = doc.styles["Normal"]
    style.font.name, style.font.size = "Calibri", Pt(11)
    h = doc.add_paragraph(); h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = h.add_run(pr["name"]); r.bold = True; r.font.size = Pt(20)
    c = doc.add_paragraph(" | ".join(contact_bits(pr)))
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph(today)
    for p in paras:
        doc.add_paragraph(p)
    doc.save(folder / f"{name}.docx")

    checks = check_letter(text, pages)
    d = job.setdefault("drafts", {})
    d.update({"cover_letter_pdf": f"{job_id}/{name}.pdf", "cover_letter_docx": f"{job_id}/{name}.docx",
              "letter_checks": checks})
    path.write_text(json.dumps(job, indent=2, ensure_ascii=False), encoding="utf-8")
    return checks


def build_master():
    pages = compile_tex(MASTER, MASTER.with_suffix(".pdf"))
    return [issue("ok" if pages <= load_profile()["resume"].get("max_pages", 1) else "error",
                  f"master-resume.pdf built: {pages} page{'s' if pages != 1 else ''}")]


if __name__ == "__main__":
    args = sys.argv[1:]
    if args == ["master"]:
        results = {"master": build_master()}
    elif len(args) == 2 and args[0] in ("resume", "letter", "all", "check"):
        what, target = args
        results = {}
        if what == "check":
            results["resume"] = build_resume_checks_only(target)
        if what in ("resume", "all"):
            results["resume"] = build_resume(target)
        if what in ("letter", "all"):
            results["letter"] = build_letter(target)
    else:
        sys.exit(__doc__)
    for k, checks in results.items():
        print(f"== {k} ==")
        for c in checks:
            print(f"  [{c['level'].upper()}] {c['text']}")
